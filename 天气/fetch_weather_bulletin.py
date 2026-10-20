# -*- coding: utf-8 -*-
"""
抓取中央气象台《天气公报》页面全部正文。

页面: https://www.nmc.cn/publish/weather-bulletin/index.htm
范围: 从标题“天气公报”到落款“制作：xxx 签发：xxx”之间的全部内容
      (含大字标题、摘要、各章节文字以及降水量/大风等预报图和图注)。

输出:
  - 天气公报_YYYYMMDD_HH时.md  (Markdown, 图片以链接形式引用)
  - 天气公报_YYYYMMDD_HH时.docx (Word, 预报图下载后嵌入; 需安装 python-docx,
    未安装时自动跳过 Word 输出)

用法:
  python fetch_weather_bulletin.py
"""

import io
import os
import re
import sys
from html.parser import HTMLParser

import requests

URL = "https://www.nmc.cn/publish/weather-bulletin/index.htm"
HEADERS = {
    "User-Agent": ("Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
                   "(KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36"),
    "Accept-Language": "zh-CN,zh;q=0.9",
}
OUT_DIR = os.path.dirname(os.path.abspath(__file__))

SECTION_RE = re.compile(r"^[一二三四五六七八九十]+、")   # “一、天气实况”这类章节标题
DATETIME_RE = re.compile(r"(\d{4})\s*年\s*(\d{1,2})\s*月\s*(\d{1,2})\s*日\s*(\d{1,2})\s*时")

PROXY_URL = os.getenv("PROXY_URL", "")
proxies = {"http": PROXY_URL, "https": PROXY_URL} if PROXY_URL else {}


class BulletinParser(HTMLParser):
    """从页面中截取 <div id=text> 块，解析出标题、署名行和正文段落。"""

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.in_text = False
        self.depth = 0
        self.zone = None            # None / 'title' / 'author' / 'writing'
        self.zone_stack = []
        self._skip = 0              # script/style 内不取文本
        self._runs = []
        self._imgs = []
        self.title = []
        self.author = []
        self.paragraphs = []        # [{'text': str, 'imgs': [url, ...]}, ...]

    # -- 辅助 --
    def _flush_to_zone(self):
        text = "".join(self._runs)
        text = re.sub(r"[\s\u3000]+", " ", text).strip()
        if self.zone == "title" and text:
            self.title.append(text)
        elif self.zone == "author" and text:
            self.author.append(text)
        elif self.zone == "writing" and (text or self._imgs):
            self.paragraphs.append({"text": text, "imgs": self._imgs})
        self._runs, self._imgs = [], []

    # -- 解析 --
    def handle_starttag(self, tag, attrs):
        if tag in ("script", "style"):
            self._skip += 1
            return
        if not self.in_text:
            if tag == "div" and ("id", "text") in attrs:
                self.in_text, self.depth, self.zone = True, 1, None
            return
        adict = dict(attrs)
        if tag == "div":
            self.zone_stack.append(self.zone)
            cls = adict.get("class", "")
            if self.zone is None:
                names = cls.split()
                if "title" in names:
                    self.zone = "title"
                elif "author" in names:
                    self.zone = "author"
                elif "writing" in names:
                    self.zone = "writing"
            self.depth += 1
        elif self.zone == "writing":
            if tag == "p":
                self._flush_to_zone()      # 段落间清空上一段残留
            elif tag == "img" and adict.get("src"):
                self._imgs.append(adict["src"])

    def handle_endtag(self, tag):
        if self._skip and tag in ("script", "style"):
            self._skip -= 1
            return
        if not self.in_text or tag != "div":
            return
        if self.zone in ("title", "author"):
            self._flush_to_zone()
        elif self.zone == "writing":
            self._flush_to_zone()
        self.zone = self.zone_stack.pop() if self.zone_stack else None
        self.depth -= 1
        if self.depth <= 0:
            self.in_text = False

    def handle_data(self, data):
        if self._skip or not self.in_text or self.zone is None:
            return
        self._runs.append(data)


def http_get(url):
    """先直连请求，失败后自动改走 proxies 代理重试。"""
    attempts = [None] + ([proxies] if proxies else [])
    last_err = None
    for px in attempts:
        label = "直连" if px is None else "代理"
        try:
            resp = requests.get(url, headers=HEADERS, proxies=px, timeout=30)
            resp.raise_for_status()
            return resp
        except requests.exceptions.RequestException as e:
            last_err = e
            print(f"  [warn] {label}请求失败: {e.__class__.__name__}，尝试下一种方式…")
    raise last_err


def fetch_html(url):
    resp = http_get(url)
    resp.encoding = "utf-8"
    return resp.text


def parse_bulletin(html):
    parser = BulletinParser()
    parser.feed(html)
    parser.close()
    title = "".join(parser.title)
    title = re.sub(r"\s+", "", title) or "天气公报"
    author = " ".join(parser.author).strip()
    author = re.sub(r"\s+", " ", author)
    author = re.sub(r"(?<!\s)签发：", "  签发：", author)   # 落款粘连处补空格

    paras = parser.paragraphs
    # 只保留到落款“制作：… 签发：…”所在段落为止
    end = max((i for i, p in enumerate(paras) if "签发：" in p["text"]), default=len(paras) - 1)
    paras = paras[: end + 1]
    return title, author, paras


def build_markdown(title, author, paras):
    lines = [f"# {title}", "", author, "", "---", ""]
    sig = next((i for i, p in enumerate(paras) if "签发：" in p["text"]), None)
    for i, p in enumerate(paras):
        text, imgs = p["text"], p["imgs"]
        if i == sig - 1 and imgs and not text:
            # 落款上方的红色分隔线是一张图片，Markdown 中无法正常显示，跳过
            # (Word 版仍保留该分隔线)
            continue
        if "签发：" in text:                       # 落款行粘连处补空格
            text = re.sub(r"(?<!\s)签发：", "  签发：", text)
        for url in imgs:
            lines.append(f"![{title}预报图]({url})")
            lines.append("")
        if not text:
            continue
        if "签发：" in text:                       # 文末落款
            lines.append(f"**{text}**")
        elif SECTION_RE.match(text):               # “一、天气实况”等章节
            lines.append(f"## {text}")
        elif paras.index(p) < 3 or text.startswith("摘要"):  # 大字标题/摘要行
            lines.append(f"**{text}**")
        elif re.match(r"^图\d+", text):            # 图注
            lines.append(f"*{text}*")
        else:
            lines.append(text)
        lines.append("")
    return "\n".join(lines).rstrip() + "\n"


def save_docx(path, title, author, paras):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
    style.font._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")

    h = doc.add_paragraph()
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = h.add_run(title)
    run.bold = True
    run.font.size = Pt(18)

    a = doc.add_paragraph()
    a.alignment = WD_ALIGN_PARAGRAPH.CENTER
    a.add_run(author)

    seen_head = 0
    for p in paras:
        text, imgs = p["text"], p["imgs"]
        if "签发：" in text:                       # 落款行粘连处补空格
            text = re.sub(r"(?<!\s)签发：", "  签发：", text)
        for url in imgs:
            try:
                data = http_get(url).content
                tmp = os.path.join(OUT_DIR, "_tmp_img")
                with open(tmp, "wb") as f:
                    f.write(data)
                doc.add_picture(tmp, width=Inches(6))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                os.remove(tmp)
            except Exception as e:
                print(f"  [warn] 图片下载失败 {url}: {e}")
        if not text:
            continue
        para = doc.add_paragraph()
        run = para.add_run(text)
        if "签发：" in text:                       # 落款
            para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            run.bold = True
        elif SECTION_RE.match(text):               # 章节标题
            run.bold = True
            run.font.size = Pt(14)
        elif text.startswith("摘要") or seen_head < 2:   # 大字标题/摘要
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run.bold = True
            if not text.startswith("摘要"):
                seen_head += 1
        elif re.match(r"^图\d+", text):            # 图注
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run.font.size = Pt(10)
        else:
            para.paragraph_format.first_line_indent = Pt(24)

    try:
        doc.save(path)
        return path
    except PermissionError:
        # 目标文件被占用（通常是在 Word/WPS 中打开着），自动换个名字另存
        alt = re.sub(r"\.docx$", "_new.docx", path)
        doc.save(alt)
        print(f"  [warn] {os.path.basename(path)} 被占用（可能正在 Word 中打开），已另存为 {os.path.basename(alt)}")
        return alt


def main():
    sys.stdout.reconfigure(errors="replace")   # 兼容 GBK 控制台，避免中文打印报错
    print(f"抓取页面: {URL}")
    html = fetch_html(URL)
    title, author, paras = parse_bulletin(html)
    if not paras:
        print("未解析到正文，页面结构可能已变化。")
        sys.exit(1)

    m = DATETIME_RE.search(author)
    stamp = "{}{:02d}{:02d}_{:02d}时".format(*map(int, m.groups())) if m else \
        __import__("datetime").datetime.now().strftime("%Y%m%d_%H时")
    base = os.path.join(OUT_DIR, f"{title}_{stamp}")

    md_path = base + ".md"
    with io.open(md_path, "w", encoding="utf-8") as f:
        f.write(build_markdown(title, author, paras))
    print(f"已保存 Markdown: {md_path}")

    try:
        import docx  # noqa: F401
    except ImportError:
        print("未安装 python-docx，跳过 Word 输出 (pip install python-docx 后重试可生成 docx)")
        return
    docx_path = save_docx(base + ".docx", title, author, paras)
    print(f"已保存 Word: {docx_path}")


if __name__ == "__main__":
    main()
